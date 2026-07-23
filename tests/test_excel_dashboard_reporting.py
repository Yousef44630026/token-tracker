"""Optional Excel dashboard preserves accounting grain and request-level latency."""

from __future__ import annotations

import gzip
import importlib.util
import json
import os
import shutil
import sys
import uuid
from pathlib import Path

if importlib.util.find_spec("pandas") is None:
    print("[SKIP] test_excel_dashboard_reporting: install .[reporting] to run")
    raise SystemExit(0)

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from tests._harness import make_checker  # noqa: E402
from tracker.models.enums import Additivity, DataQualityFlag, PrecisionLevel, TokenType, UsageSource  # noqa: E402
from tracker.models.token_event import TokenEvent  # noqa: E402
from tracker.models.token_quantity import TokenQuantity  # noqa: E402
from tracker.observability.observation import Observation  # noqa: E402
from tracker.reporting.excel_dashboard import (  # noqa: E402
    LoadedEvent,
    _billing_allocations,
    _find_price,
    build_data_frame,
    build_summary_frames,
    load_jsonl_events,
    load_prices,
    write_dashboard,
)

check = make_checker()
root = Path(f".test_excel_dashboard_{uuid.uuid4().hex}").resolve()
data_dir = root / "data"
data_dir.mkdir(parents=True)


def quantity(token_type, value, additivity=Additivity.TOTAL_CONTRIBUTING, *, subtotal_of=None):
    return TokenQuantity(
        token_type=token_type,
        quantity=value,
        precision_level=PrecisionLevel.EXACT,
        usage_source=UsageSource.PROVIDER_RESPONSE,
        additivity=additivity,
        subtotal_of=subtotal_of,
    )


partial = TokenEvent(
    event_id="partial",
    request_correlation_id="request-1",
    trace_id="trace-1",
    span_id="span-1",
    workflow="support",
    provider="azure_openai",
    model="model-a",
    api_surface="responses",
    quantities=[
        TokenQuantity(
            token_type=TokenType.OUTPUT,
            quantity=40,
            precision_level=PrecisionLevel.ESTIMATE,
            usage_source=UsageSource.PARTIAL_STREAM_TOKENIZER,
            additivity=Additivity.TOTAL_CONTRIBUTING,
        )
    ],
    data_quality_flags=[DataQualityFlag.PARTIAL_STREAM_ESTIMATE.value],
    timestamp="2026-07-01T09:59:00Z",
    observation=Observation(authoritative=True, status="incomplete", duration_ms=50),
)
final_one = TokenEvent(
    event_id="final-1",
    request_correlation_id="request-1",
    trace_id="trace-1",
    span_id="span-1",
    workflow="support",
    provider="azure_openai",
    model="model-a",
    api_surface="responses",
    quantities=[
        quantity(TokenType.INPUT, 100),
        quantity(TokenType.CACHED_INPUT, 20, Additivity.SUBTOTAL_OF, subtotal_of="input"),
        quantity(TokenType.OUTPUT, 50),
    ],
    provider_total_tokens=150,
    timestamp="2026-07-01T10:00:00Z",
    observation=Observation(
        authoritative=True,
        status="success",
        duration_ms=100,
        time_to_first_token_ms=20,
        extra={"source": "azure_smoke", "source_event_id": "source-1"},
    ),
)
for final_one_quantity in final_one.quantities:
    final_one_quantity.metadata["azure_deployment"] = "deployment-a"
final_two = TokenEvent(
    event_id="final-2",
    request_correlation_id="request-2",
    trace_id="trace-2",
    span_id="span-2",
    workflow="summarization",
    provider="azure_openai",
    model="model-b",
    api_surface="responses",
    quantities=[quantity(TokenType.INPUT, 200), quantity(TokenType.OUTPUT, 100)],
    provider_total_tokens=300,
    timestamp="2026-07-02T11:00:00Z",
    observation=Observation(authoritative=True, status="success", duration_ms=300),
)

events_path = data_dir / "events.jsonl"
with events_path.open("w", encoding="utf-8") as handle:
    for event in (partial, final_one, final_two):
        handle.write(json.dumps(event.to_dict()) + "\n")
    handle.write("{not-json\n")
    handle.write(json.dumps({"not": "an event"}) + "\n")

prices_path = root / "prices.csv"
pd.DataFrame(
    [
        ("azure_openai", "model-a", "input", 1.0, "USD", "2026-01-01", ""),
        ("azure_openai", "model-a", "cached_input", 0.2, "USD", "2026-01-01", ""),
        ("azure_openai", "model-a", "output", 2.0, "USD", "2026-01-01", ""),
        ("azure_openai", "model-b", "input", 0.5, "USD", "2026-01-01", ""),
        ("azure_openai", "model-b", "output", 1.5, "USD", "2026-01-01", ""),
    ],
    columns=("provider", "model", "token_type", "price_per_million_tokens", "currency", "effective_from", "effective_to"),
).to_csv(prices_path, index=False)

price_columns = (
    "provider",
    "model",
    "token_type",
    "price_per_million_tokens",
    "currency",
    "effective_from",
    "effective_to",
)


def rejects_price_table(name, rows):
    candidate = root / f"prices-{name}.csv"
    pd.DataFrame(rows, columns=price_columns).to_csv(candidate, index=False)
    try:
        load_prices(candidate)
    except ValueError:
        return True
    return False


check(
    rejects_price_table("bad-date", [("azure_openai", "model-a", "input", 1, "USD", "tomorrow", "")]),
    "invalid effective dates are rejected instead of becoming unbounded",
)
check(
    rejects_price_table("inverted", [("azure_openai", "model-a", "input", 1, "USD", "2026-02-01", "2026-01-01")]),
    "inverted effective ranges are rejected",
)
check(
    rejects_price_table(
        "overlap",
        [
            ("azure_openai", "model-a", "input", 1, "USD", "2026-01-01", "2026-06-30"),
            ("azure_openai", "model-a", "input", 2, "USD", "2026-06-30", "2026-12-31"),
        ],
    ),
    "overlapping prices for one selector are rejected",
)
check(
    rejects_price_table("blank-currency", [("azure_openai", "model-a", "input", 1, "", "", "")]),
    "blank pricing dimensions are rejected",
)
check(
    rejects_price_table("unknown-type", [("azure_openai", "model-a", "mystery_tokens", 1, "USD", "", "")]),
    "unknown token types cannot silently miss every event",
)

loaded_prices = load_prices(prices_path)
check(
    _find_price(
        loaded_prices,
        provider="azure_openai",
        model="model-a",
        token_type="input",
        event_date=None,
    )
    is None,
    "an event without a timestamp never receives an effective-dated price",
)

ambiguous_prices_path = root / "prices-ambiguous-match.csv"
pd.DataFrame(
    [
        ("azure_openai", "*", "input", 1, "USD", "", ""),
        ("*", "model-a", "input", 2, "USD", "", ""),
    ],
    columns=price_columns,
).to_csv(ambiguous_prices_path, index=False)
try:
    _find_price(
        load_prices(ambiguous_prices_path),
        provider="azure_openai",
        model="model-a",
        token_type="input",
        event_date=None,
    )
except ValueError:
    ambiguous_match_rejected = True
else:
    ambiguous_match_rejected = False
check(ambiguous_match_rejected, "equally specific price selectors are rejected instead of chosen arbitrarily")

unverified_event = TokenEvent(
    event_id="unverified-price",
    request_correlation_id="unverified-price",
    trace_id="trace-price",
    span_id="span-price",
    provider="azure_openai",
    model="model-a",
    timestamp="2026-07-01T00:00:00Z",
    quantities=[quantity(TokenType.INPUT, 100, Additivity.UNVERIFIED)],
    observation=Observation(authoritative=True, status="success"),
)
unverified_frame = build_data_frame(
    [LoadedEvent(unverified_event, "memory", 1, 1)],
    loaded_prices,
)
check(unverified_frame["derived_cost"].isna().all(), "unverified additivity never produces a plausible cost")
check(unverified_frame["cost_quality"].iloc[0] == "unverified_additivity", "unverified cost exclusion carries a reason")
check(unverified_frame["unverified_tokens_active"].sum() == 100, "dashboard quantifies known unverified token magnitude")

unknown_event = TokenEvent(
    event_id="unknown-price",
    request_correlation_id="unknown-price",
    trace_id="trace-price",
    span_id="span-price-unknown",
    provider="azure_openai",
    model="model-a",
    timestamp="2026-07-01T00:00:00Z",
    quantities=[
        TokenQuantity(
            token_type=TokenType.INPUT,
            quantity=None,
            precision_level=PrecisionLevel.UNKNOWN,
            usage_source=UsageSource.PROVIDER_RESPONSE,
            additivity=Additivity.TOTAL_CONTRIBUTING,
        )
    ],
    observation=Observation(authoritative=True, status="complete"),
)
unknown_data = build_data_frame([LoadedEvent(unknown_event, "memory", 1, 1)], loaded_prices)
unknown_summary = build_summary_frames(unknown_data)
check(unknown_summary["total_cost"] is None, "unknown token magnitude keeps derived cost unknown")
check(unknown_summary["quality"]["coverage_status"] == "missing", "unknown-only data cannot claim complete coverage")

cross_cutting_event = TokenEvent(
    event_id="cross-cutting-subtotals",
    request_correlation_id="cross-cutting-subtotals",
    trace_id="trace-price",
    span_id="span-price",
    quantities=[
        quantity(TokenType.INPUT, 100),
        quantity(TokenType.CACHED_INPUT, 20, Additivity.SUBTOTAL_OF, subtotal_of="input"),
        quantity(TokenType.IMAGE_INPUT, 30, Additivity.SUBTOTAL_OF, subtotal_of="input"),
    ],
    observation=Observation(authoritative=True, status="success"),
)
cross_allocations, cross_issues = _billing_allocations(cross_cutting_event)
check(
    all(value is None for value in cross_allocations)
    and set(cross_issues) == {"ambiguous_subtotal_overlap"},
    "cross-cutting cache/modality subtotals fail closed instead of double-subtracting",
)

events, report = load_jsonl_events(data_dir)
check(Path(f"{events_path}.lock").exists(), "dashboard reads JSONL under the repository interprocess lock")
check(report.valid_events == 3, "three valid events survive JSONL validation")
check(report.malformed_lines == 1, "malformed JSONL is logged and skipped")
check(report.schema_invalid_lines == 1, "schema-invalid JSON object is logged and skipped")
check(next(item.event for item in events if item.event.event_id == "partial").superseded, "final usage supersedes partial stream")

identity_dir = root / "identity"
identity_dir.mkdir()
canonical_duplicate = TokenEvent(
    event_id="duplicate-identity",
    request_correlation_id="duplicate-identity-request",
    trace_id="duplicate-identity-trace",
    span_id="duplicate-identity-span",
    quantities=[quantity(TokenType.INPUT, 10)],
    timestamp="2026-07-01T00:00:00Z",
    observation=Observation(authoritative=True, status="complete"),
)
contradictory_duplicate = TokenEvent.from_dict(canonical_duplicate.to_dict())
contradictory_duplicate.quantities[0].quantity = 999
with (identity_dir / "events.jsonl").open("w", encoding="utf-8") as handle:
    handle.write(json.dumps(canonical_duplicate.to_dict()) + "\n")
    handle.write(json.dumps(contradictory_duplicate.to_dict()) + "\n")
identity_events, identity_report = load_jsonl_events(identity_dir)
check(identity_report.duplicate_event_ids == 1, "contradictory duplicate event identities are reported")
check(
    identity_events[0].event.quantities[0].quantity == 10,
    "Excel loading preserves the core first-event-id-wins identity rule",
)

archive_dir = data_dir / "events.jsonl.archive"
archive_dir.mkdir()
with gzip.open(archive_dir / "segment.jsonl.gz", "wt", encoding="utf-8") as handle:
    handle.write(json.dumps(final_two.to_dict()) + "\n")
archived_events, archived_report = load_jsonl_events(data_dir)
check(archived_report.files_read == 2, "dashboard discovers active and archived JSONL segments")
check(archived_report.duplicate_event_ids == 1, "dashboard deduplicates archive/active crash overlap")
check(len(archived_events) == 3, "dashboard archive discovery does not double count")

data = build_data_frame(events, loaded_prices)
summaries = build_summary_frames(data)
check(data["event_contributing_tokens_once"].sum() == 450, "event-grain safe total is never repeated at quantity grain")
check(data["exact_tokens_active"].sum() == 450, "dashboard exposes exact contributing-token magnitude")
check(data["estimated_tokens_active"].sum() == 0, "superseded estimates do not pollute active estimate magnitude")
check(data["event_count_once"].sum() == 3, "event count appears exactly once at quantity grain")
check(data["event_authoritative_once"].sum() == 3, "authoritative event count appears exactly once")
check(data["superseded_event_once"].sum() == 1, "superseded event KPI is additive and event-grain safe")
check(data["active_quality_flagged_event_once"].sum() == 0, "superseded quality flags do not pollute active anomalies")
check(data["mismatch_event_once"].sum() == 0, "active mismatch event KPI is additive")
check(
    set(data.loc[(data["model"] == "model-a") & ~data["event_superseded"], "deployment"])
    == {"deployment-a"},
    "Azure deployment metadata reaches the flat dashboard source",
)
check(set(data.loc[data["provider"] == "azure_openai", "cloud_provider"]) == {"azure"}, "cloud attribution is explicit")
check(data["request_count_once"].sum() == 2, "request count appears exactly once per correlation id")
check(data["request_latency_observation_once"].sum() == 2, "latency denominator counts requests, not quantities")
check(data.loc[data["request_count_once"] == 1, "request_latency_ms"].mean() == 200, "latency average uses one row per request")
check(set(data.loc[data["event_id"] == "final-1", "source_kind"]) == {"azure_smoke"}, "telemetry source is flattened")
check(
    data.loc[data["event_id"] == "final-1", "observation_json"].notna().sum() == 1,
    "event observation JSON is serialized once instead of repeated at quantity grain",
)

model_a = data[(data["model"] == "model-a") & ~data["event_superseded"]]
check(model_a.loc[model_a["token_type"] == "input", "billing_tokens"].iloc[0] == 80, "input billing removes cached subtotal")
check(model_a.loc[model_a["token_type"] == "cached_input", "billing_tokens"].iloc[0] == 20, "cache subtotal has its own billing rate")
check(abs(float(data["derived_cost"].sum()) - 0.000434) < 1e-12, "derived cost uses allocated tokens without double counting")
check(data["billable_tokens_for_coverage"].sum() == 450, "pricing coverage denominator excludes superseded usage")
check(data["priced_billing_tokens"].sum() == 450, "pricing coverage numerator carries priced token magnitude")
check(data["cache_read_tokens_active"].sum() == 20, "cache KPI is additive without event-level boolean criteria")
check(data["unknown_quantity_active"].sum() == 0, "unknown quantity KPI excludes superseded estimates")
check(abs(float(summaries["pricing_coverage"]) - 1.0) < 1e-12, "pricing coverage is complete")
check(summaries["quality"]["latency_coverage"] == 1.0, "dashboard quantifies request-level latency coverage")
check(
    summaries["quality"]["instrumented_latency_coverage"] == 1.0,
    "instrumentable requests expose a separate latency coverage",
)
check(summaries["quality"]["latency_applicability"] == 1.0, "live request latency is fully applicable")
check(summaries["quality"]["provider_total_coverage"] == 1.0, "dashboard quantifies provider-total coverage")
check(summaries["quality"]["quality_status"] == "clean", "fully covered active data receives a clean runtime status")
check(summaries["quality"]["coverage_status"] == "complete", "complete price and latency coverage is explicit")
check(not summaries["provider_summary"].empty, "provider runtime quality summary is available")
check(not summaries["source_summary"].empty, "source provenance summary is available")

under_attributed_event = TokenEvent(
    event_id="under-attributed-dashboard",
    request_correlation_id="under-attributed-dashboard",
    trace_id="trace-quality",
    span_id="span-quality",
    provider="azure_openai",
    model="model-a",
    timestamp="2026-07-01T00:00:00Z",
    quantities=[quantity(TokenType.INPUT, 100)],
    provider_total_tokens=120,
    observation=Observation(authoritative=True, status="success", duration_ms=10),
)
under_summary = build_summary_frames(
    build_data_frame([LoadedEvent(under_attributed_event, "memory", 1, 1)], loaded_prices)
)
check(
    under_summary["quality"]["quality_status"] == "warning",
    "provider under-attribution prevents a falsely clean dashboard status",
)

unpriced_data = build_data_frame(events, load_prices(None))
unpriced_summaries = build_summary_frames(unpriced_data)
check(unpriced_summaries["total_cost"] is None, "an entirely unpriced workload has unknown cost, never zero")
check(unpriced_summaries["pricing_coverage"] == 0.0, "unpriced token magnitude has zero pricing coverage")
check(unpriced_summaries["quality"]["quality_status"] == "clean", "missing prices do not imply corrupt token data")
check(unpriced_summaries["quality"]["coverage_status"] == "partial", "missing prices remain a separate coverage gap")
check(unpriced_summaries["cost_by_day"].empty, "unpriced workload does not fabricate zero-valued cost trends")
check(
    unpriced_summaries["use_cases"]["derived_cost"].isna().all(),
    "unpriced use-case costs remain unknown",
)

local_import = TokenEvent(
    event_id="claude-local-import",
    request_correlation_id="claude-local-import-request",
    trace_id="claude-local-import-trace",
    span_id="claude-local-import-span",
    provider="anthropic",
    model="claude-test",
    api_surface="messages",
    quantities=[quantity(TokenType.INPUT, 50)],
    data_quality_flags=[DataQualityFlag.CLAUDE_CODE_LOCAL_USAGE.value],
    timestamp="2026-07-01T00:01:00Z",
    observation=Observation(authoritative=True, status="complete"),
)
mixed_latency = build_summary_frames(
    build_data_frame(
        [LoadedEvent(final_two, "memory", 1, 1), LoadedEvent(local_import, "memory", 2, 2)],
        loaded_prices,
    )
)["quality"]
check(mixed_latency["latency_coverage"] == 0.5, "overall latency presence still exposes local-log gaps")
check(
    mixed_latency["instrumented_latency_coverage"] == 1.0,
    "unobservable local imports do not poison instrumented latency coverage",
)
check(mixed_latency["latency_applicability"] == 0.5, "latency applicability quantifies the excluded source mix")

output = root / "dashboard.xlsx"
write_dashboard(data, summaries, report, output)
try:
    write_dashboard(data, summaries, report, root / "oversized.xlsx", max_data_rows=len(data) - 1)
except ValueError as exc:
    dashboard_limit_rejected = "safety limit" in str(exc)
else:
    dashboard_limit_rejected = False
check(dashboard_limit_rejected, "dashboard fails before producing an intentionally oversized workbook")
workbook = load_workbook(output, data_only=False)
visible_sheets = [sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
check(
    visible_sheets
    == ["Data", "Dashboard", "Data Quality", "Provider Readiness", "Coûts", "Tokens & Latence", "Use cases"],
    "workbook exposes dedicated runtime-quality and provider-certification sheets",
)
check(workbook["_Lists"].sheet_state == "veryHidden", "filter support lists are not exposed as a reporting sheet")
check("DataTable" in workbook["Data"].tables, "Data is an Excel table suitable for pivots")
dashboard = workbook["Dashboard"]
check(workbook.active.title == "Dashboard", "interactive dashboard opens first")
check(dashboard["A1"].value == "Multi-cloud LLM token observability", "dashboard title matches Azure, Google, and AWS scope")
check(len(dashboard.data_validations.dataValidation) == 7, "dashboard has five dimension and two date selectors")
check(
    {"DashboardProviders", "DashboardModels", "DashboardDeployments", "DashboardEnvironments", "DashboardUseCases"}
    <= set(workbook.defined_names),
    "filter dropdowns use workbook ranges instead of the 255-character inline-list limit",
)
check(dashboard["B5"].value == "All" and dashboard["L5"].value is not None, "dashboard filters have safe defaults")
check(
    dashboard["A11"].value == "Known exact token share"
    and dashboard["C11"].value == "Estimated tokens"
    and dashboard["G11"].value == "Under-attributed"
    and dashboard["K11"].value == "Schema drift events",
    "integrity magnitudes are visible in the dashboard quality row",
)
check(dashboard["M11"].value == "Correlation risks", "correlation-id uncertainty is a headline KPI")
check(
    dashboard["K7"].value == "Pricing coverage" and dashboard["M7"].value == "Latency coverage",
    "pricing and latency coverage are explicit headline KPIs",
)
check("COST COVERAGE INCOMPLETE" in dashboard["A15"].value, "dashboard includes a filter-aware readiness banner")
check("ISNUMBER(K8)" in dashboard["A15"].value, "readiness banner handles non-applicable coverage without formula errors")
check(
    isinstance(dashboard["A8"].value, str) and "DataTable[derived_cost]" in dashboard["A8"].value,
    "cost KPI is an Excel formula over the Data table",
)
check(
    "DataTable[unknown_quantity_active]" in dashboard["A8"].value,
    "interactive cost fails closed when the selected data contains an unknown quantity",
)
check(
    dashboard["B51"].value == "2026-07-01"
    and isinstance(dashboard["C51"].value, str)
    and "AND(" in dashboard["C51"].value,
    "daily chart series honor the selected date range",
)
check(len(dashboard._charts) == 4, "dashboard has four filter-driven native charts")
check(
    [series.tx.v for series in dashboard._charts[2].series] == ["Average latency"],
    "latency chart explicitly names the observed average series",
)
check(dashboard._charts[2].display_blanks == "gap", "unknown P95 values are gaps, never fabricated zeros")
check({"DashboardDaily", "DashboardModelSummary"} <= set(dashboard.tables), "chart helper ranges remain auditable")
check(len(workbook["Coûts"]._charts) == 2, "cost sheet has two native charts")
check(len(workbook["Tokens & Latence"]._charts) == 2, "tokens and latency sheet has two native charts")
check(workbook["Coûts"]["A9"].value == "2026-07-01", "static chart dates use readable ISO labels")
check(
    [series.tx.v for series in workbook["Tokens & Latence"]._charts[0].series] == ["input", "output"],
    "token chart series have semantic names instead of generic labels",
)
check(
    workbook["Tokens & Latence"]._charts[1].display_blanks == "gap",
    "per-model latency charts never turn missing observations into zero milliseconds",
)
check(len(workbook["Use cases"]._charts) == 1, "use-case sheet has a native pie chart")
quality_sheet = workbook["Data Quality"]
check(quality_sheet["A1"].value == "Data quality and provenance", "quality sheet has an audit-oriented title")
check(quality_sheet["A5"].value == "clean", "quality sheet publishes the runtime quality status")
check(quality_sheet["K5"].value == "complete", "quality sheet separates operational coverage status")
check("ProviderRuntimeQuality" in quality_sheet.tables, "quality sheet includes provider-level coverage")
check("SourceProvenance" in quality_sheet.tables, "quality sheet explains where token totals came from")
readiness_sheet = workbook["Provider Readiness"]
check(readiness_sheet["A1"].value == "Provider readiness", "provider evidence is visible in the workbook")
check("ProviderSurfaceCertification" in readiness_sheet.tables, "surface-level certification remains auditable")
check("ProviderCapabilityCertification" in readiness_sheet.tables, "capability-level certification remains auditable")
check(not any("derived_cost" in event.to_dict() for event in (partial, final_one, final_two)), "derived cost never enters stored events")

workbook.close()
demo_output = os.environ.get("TRACKER_DASHBOARD_DEMO_OUTPUT")
if demo_output:
    Path(demo_output).resolve().parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(output, demo_output)
shutil.rmtree(root, ignore_errors=True)
sys.exit(check.report("RESULT test_excel_dashboard_reporting"))
